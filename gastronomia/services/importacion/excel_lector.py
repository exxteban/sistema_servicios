"""Lector tolerante de planillas de ventas.

No depende de posiciones fijas: busca los encabezados en cada hoja, decide que
tipo de hoja es (detalle de ventas, totales por dia, egresos) y arma bloques de
venta por dia separados por filas "Total" o por un cambio de fecha.
"""
from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal

from openpyxl import load_workbook

from gastronomia.services.importacion.fechas import parse_fecha
from gastronomia.services.importacion.texto import normalizar, parse_cantidad, parse_monto


SINONIMOS = {
    'fecha': ('fecha', 'dia', 'fechas'),
    'cantidad': ('cantidad', 'cant', 'unidades', 'unid', 'qty', 'cantidades'),
    'producto': ('mercaderia', 'mercaderias', 'producto', 'productos', 'detalle', 'descripcion',
                 'item', 'articulo', 'plato', 'pedido', 'concepto venta'),
    'monto': ('monto', 'total', 'importe', 'valor', 'subtotal', 'precio', 'ingreso', 'montos'),
}
PALABRAS_EGRESO = ('egreso', 'gasto', 'razon social', 'proveedor', 'salario', 'compra')
FILAS_BUSQUEDA_ENCABEZADO = 30


@dataclass
class FilaVenta:
    fila: int
    cantidad: int
    nombre: str
    monto: Decimal | None


@dataclass
class Bloque:
    hoja: str
    fila_inicio: int
    fila_fin: int = 0
    fecha_texto: str = ''
    fecha: date | None = None
    total_excel: Decimal | None = None
    filas: list[FilaVenta] = field(default_factory=list)

    @property
    def suma(self) -> Decimal:
        return sum((item.monto or Decimal('0')) for item in self.filas) or Decimal('0')


@dataclass
class TotalDiario:
    hoja: str
    fila: int
    fecha: date
    monto: Decimal


@dataclass
class LecturaExcel:
    bloques: list[Bloque] = field(default_factory=list)
    totales_diarios: list[TotalDiario] = field(default_factory=list)
    hojas: list[dict] = field(default_factory=list)
    anio_defecto: int = 0


def leer_excel(ruta: str) -> LecturaExcel:
    libro = load_workbook(ruta, data_only=True, read_only=True)
    try:
        hojas = [(hoja.title, [list(fila) for fila in hoja.iter_rows(values_only=True)]) for hoja in libro.worksheets]
    finally:
        libro.close()
    lectura = LecturaExcel(anio_defecto=_anio_mas_comun(hojas))
    for nombre, filas in hojas:
        _leer_hoja(lectura, nombre, filas)
    return lectura


def _leer_hoja(lectura: LecturaExcel, nombre: str, filas: list[list]) -> None:
    encabezado = _buscar_encabezado(filas)
    tipo = _tipo_hoja(nombre, encabezado, filas, lectura.anio_defecto)
    info = {'nombre': nombre, 'tipo': tipo, 'filas': len(filas)}
    if tipo == 'ventas':
        fila_header, columnas = encabezado
        bloques = _leer_bloques(nombre, filas, fila_header, columnas, lectura.anio_defecto)
        lectura.bloques.extend(bloques)
        info.update({'bloques': len(bloques), 'columnas': _columnas_legibles(columnas)})
    elif tipo == 'totales':
        totales = _leer_totales(nombre, filas, lectura.anio_defecto)
        lectura.totales_diarios.extend(totales)
        info['dias'] = len(totales)
    lectura.hojas.append(info)


def _buscar_encabezado(filas: list[list]) -> tuple[int, dict] | None:
    mejor = None
    for indice, fila in enumerate(filas[:FILAS_BUSQUEDA_ENCABEZADO]):
        columnas = {}
        for col, valor in enumerate(fila):
            rol = _rol_encabezado(valor)
            if rol and rol not in columnas:
                columnas[rol] = col
        if 'monto' in columnas and ('producto' in columnas or 'fecha' in columnas):
            if not mejor or len(columnas) > len(mejor[1]):
                mejor = (indice, columnas)
    return mejor


def _rol_encabezado(valor) -> str | None:
    texto = normalizar(valor)
    if not texto or len(texto) > 30:
        return None
    for rol, sinonimos in SINONIMOS.items():
        if texto in sinonimos or any(texto.startswith(f'{item} ') for item in sinonimos):
            return rol
    return None


def _tipo_hoja(nombre: str, encabezado, filas: list[list], anio: int) -> str:
    nombre_norm = normalizar(nombre)
    textos_header = ''
    if encabezado:
        textos_header = ' '.join(normalizar(valor) for valor in filas[encabezado[0]] if valor is not None)
    if any(palabra in nombre_norm or palabra in textos_header for palabra in PALABRAS_EGRESO):
        return 'egresos'
    if encabezado and 'producto' in encabezado[1]:
        return 'ventas'
    if _leer_totales(nombre, filas, anio):
        return 'totales'
    return 'vacia' if not any(any(c is not None for c in fila) for fila in filas) else 'otra'


def _leer_bloques(hoja: str, filas: list[list], fila_header: int, columnas: dict, anio: int) -> list[Bloque]:
    bloques: list[Bloque] = []
    actual: Bloque | None = None

    def cerrar(fin: int):
        nonlocal actual
        if actual and (actual.filas or actual.total_excel is not None):
            actual.fila_fin = fin
            bloques.append(actual)
        actual = None

    for indice in range(fila_header + 1, len(filas)):
        fila = filas[indice]
        numero = indice + 1
        valor_fecha = _celda(fila, columnas.get('fecha'))
        nombre = str(_celda(fila, columnas.get('producto')) or '').strip()
        monto = parse_monto(_celda(fila, columnas.get('monto')))
        if _es_total(nombre) or _es_total(valor_fecha):
            if actual is None:
                actual = Bloque(hoja=hoja, fila_inicio=numero)
            actual.total_excel = monto
            cerrar(numero)
            continue
        if valor_fecha not in (None, '') and actual and actual.filas:
            cerrar(numero - 1)
        if not nombre and monto is None:
            if valor_fecha not in (None, ''):
                actual = actual or Bloque(hoja=hoja, fila_inicio=numero)
                _asignar_fecha(actual, valor_fecha, anio)
            continue
        if actual is None:
            actual = Bloque(hoja=hoja, fila_inicio=numero)
        if valor_fecha not in (None, '') and not actual.fecha_texto:
            _asignar_fecha(actual, valor_fecha, anio)
        cantidad, nombre_limpio = _cantidad_y_nombre(_celda(fila, columnas.get('cantidad')), nombre)
        actual.filas.append(FilaVenta(fila=numero, cantidad=cantidad, nombre=nombre_limpio, monto=monto))
    cerrar(len(filas))
    return bloques


def _asignar_fecha(bloque: Bloque, valor, anio: int) -> None:
    bloque.fecha_texto = valor.strftime('%d/%m/%Y') if isinstance(valor, (datetime, date)) else str(valor).strip()
    bloque.fecha = parse_fecha(valor, anio)


def _cantidad_y_nombre(valor_cantidad, nombre: str) -> tuple[int, str]:
    cantidad = parse_cantidad(valor_cantidad)
    match = re.match(r'^\s*(\d{1,2})\s*(?:x\s*)?(?=[^\d\s])', nombre, flags=re.IGNORECASE)
    if match:
        nombre = nombre[match.end():].strip()
        cantidad = cantidad or int(match.group(1))
    return cantidad or 1, nombre


def _leer_totales(hoja: str, filas: list[list], anio: int) -> list[TotalDiario]:
    totales = []
    for indice, fila in enumerate(filas):
        celdas = [valor for valor in fila if valor not in (None, '')]
        if len(celdas) < 2 or any(_es_total(valor) for valor in celdas):
            continue
        fecha = next((parse_fecha(v, anio) for v in celdas if isinstance(v, (datetime, date))), None)
        montos = [parse_monto(v) for v in celdas if isinstance(v, (int, float, Decimal))]
        montos = [m for m in montos if m and m > 0]
        if fecha and len(montos) == 1:
            totales.append(TotalDiario(hoja=hoja, fila=indice + 1, fecha=fecha, monto=montos[0]))
    return totales


def _es_total(valor) -> bool:
    texto = normalizar(valor)
    return bool(texto) and (texto == 'total' or texto.startswith('total ') or texto == 'totales')


def _celda(fila: list, columna: int | None):
    if columna is None or columna >= len(fila):
        return None
    return fila[columna]


def _columnas_legibles(columnas: dict) -> dict:
    letras = {}
    for rol, col in columnas.items():
        letras[rol] = _letra_columna(col)
    return letras


def _letra_columna(indice: int) -> str:
    letras = ''
    indice += 1
    while indice:
        indice, resto = divmod(indice - 1, 26)
        letras = chr(65 + resto) + letras
    return letras


def _anio_mas_comun(hojas) -> int:
    anios = Counter(
        valor.year
        for _nombre, filas in hojas
        for fila in filas
        for valor in fila
        if isinstance(valor, (datetime, date))
    )
    return anios.most_common(1)[0][0] if anios else date.today().year
