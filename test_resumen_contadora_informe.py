"""Descargas del mes para el/la contador/a: PDF (completo o sólo resumen) y Excel.

Salen de los mismos cálculos que la pantalla: los números tienen que coincidir.
"""
from datetime import date, datetime
from io import BytesIO

from app import db
from app.models import Compra, Configuracion, Proveedor, Usuario, Venta
from facturacion_electronica import CLAVE_FACTURACION_ELECTRONICA_ACTIVO

from test_facturacion_electronica_proveedor_base import BaseProveedorApi

MES = '2026-09'


class TestInformeContadora(BaseProveedorApi):
    def setUp(self):
        super().setUp()
        Configuracion.establecer(CLAVE_FACTURACION_ELECTRONICA_ACTIVO, 'false')
        Configuracion.establecer('nombre_empresa', 'Pablito Cell')
        Configuracion.establecer('ruc_empresa', '80099999-1')
        db.session.commit()
        self.http = self.app.test_client()
        with self.http.session_transaction() as sesion:
            sesion['_user_id'] = str(Usuario.query.filter_by(username='admin').first().id_usuario)
            sesion['_fresh'] = True

    def _datos(self):
        for numero, estado in (('001-001-0000101', 'completada'), ('001-001-0000102', 'completada'),
                               ('001-001-0000103', 'anulada')):
            db.session.add(Venta(id_cliente=self.cliente.id_cliente, id_sesion_caja=self.sesion.id_sesion,
                                 subtotal=110000, total=110000, total_iva_10=10000, numero_factura=numero,
                                 estado=estado, fecha_venta=datetime(2026, 9, 15, 15)))
        proveedor = Proveedor(nombre='Distribuidora Sur', ruc='80012345-6')
        db.session.add(proveedor)
        db.session.flush()
        db.session.add(Compra(numero_factura='001-002-0000555', timbrado='15151515', id_proveedor=proveedor.id_proveedor,
                              id_usuario=Usuario.query.filter_by(username='admin').first().id_usuario,
                              fecha_compra=date(2026, 9, 10), subtotal=55000, total=55000,
                              total_iva_10=5000, total_iva_5=0, estado='completada'))
        db.session.commit()
        self.http.post('/contador/iva/ajustes', data={'mes': MES, 'saldo_favor_anterior': '3.000', 'retenciones': '0'})

    def _texto_pdf(self, respuesta):
        from pypdf import PdfReader

        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(respuesta.headers['Content-Type'], 'application/pdf')
        lector = PdfReader(BytesIO(respuesta.data))
        return '\n'.join(pagina.extract_text() for pagina in lector.pages), len(lector.pages)

    def test_pdf_completo_con_resumen_observaciones_y_anexos(self):
        self._datos()

        texto, paginas = self._texto_pdf(self.http.get(f'/contador/iva/pdf?mes={MES}'))

        self.assertIn('Resumen de IVA', texto)
        self.assertIn('Septiembre 2026', texto)
        self.assertIn('Pablito Cell', texto)
        self.assertIn('80099999-1', texto)
        self.assertIn('20.000', texto)                          # débito: 2 ventas con factura
        self.assertIn('12.000', texto)                          # saldo: 20.000 − 5.000 − 3.000
        self.assertIn('Pendiente para cerrar el mes', texto)
        self.assertIn('RUC con el dígito verificador mal', texto)
        self.assertIn('Anexo 1', texto)
        self.assertIn('001-001-0000101', texto)
        self.assertIn('001-002-0000555', texto)                 # compra en su anexo
        self.assertIn('15151515', texto)                        # con su timbrado
        self.assertIn('001-001-0000103', texto)                 # anulada en su anexo
        self.assertGreaterEqual(paginas, 6)                     # resumen + 5 anexos

    def test_pdf_solo_resumen_no_trae_anexos(self):
        self._datos()

        respuesta = self.http.get(f'/contador/iva/pdf?mes={MES}&anexos=0')
        texto, paginas = self._texto_pdf(respuesta)

        self.assertIn('Cuadro de IVA', texto)
        self.assertNotIn('Anexo 1', texto)
        self.assertLess(paginas, 3)
        self.assertIn('iva_2026_09_resumen.pdf', respuesta.headers['Content-Disposition'])

    def test_excel_completo_con_una_hoja_por_anexo(self):
        from openpyxl import load_workbook

        self._datos()

        respuesta = self.http.get(f'/contador/iva/excel?mes={MES}')

        self.assertEqual(respuesta.status_code, 200)
        libro = load_workbook(BytesIO(respuesta.data))
        self.assertEqual(libro.sheetnames, ['Resumen', 'Ventas con número de factura', 'Notas de crédito',
                                            'Compras con número de factura', 'Gastos con factura', 'Facturas anuladas'])
        resumen = [[c for c in fila if c is not None] for fila in libro['Resumen'].iter_rows(values_only=True)]
        self.assertIn(['Débito fiscal', 20000], resumen)
        self.assertIn(['Saldo del mes (+ a pagar / − a favor)', 12000], resumen)
        compras = list(libro['Compras con número de factura'].iter_rows(values_only=True))
        self.assertIn('Timbrado', compras[3])
        self.assertTrue(any(fila[1] == '001-002-0000555' for fila in compras))

    def test_excel_trae_todas_las_anuladas_aunque_la_pantalla_corte(self):
        """La pantalla muestra una página; el informe descargable no puede cortar."""
        from openpyxl import load_workbook
        from app.services.resumen_contadora_ventas import FILAS_ANULADAS

        cantidad = FILAS_ANULADAS + 5
        for n in range(cantidad):
            db.session.add(Venta(id_cliente=self.cliente.id_cliente, id_sesion_caja=self.sesion.id_sesion,
                                 subtotal=11000, total=11000, total_iva_10=1000,
                                 numero_factura=f'001-001-{n + 1:07d}', estado='anulada',
                                 fecha_venta=datetime(2026, 9, 15, 15)))
        db.session.commit()

        libro = load_workbook(BytesIO(self.http.get(f'/contador/iva/excel?mes={MES}').data))
        numeros = [fila[1] for fila in libro['Facturas anuladas'].iter_rows(values_only=True)
                   if isinstance(fila[1], str) and fila[1].startswith('001-001-')]
        self.assertEqual(len(numeros), cantidad)

    def test_hora_de_generado_es_la_del_negocio(self):
        from app.services.resumen_contadora_informe import informe_del_mes

        self.assertIsNotNone(informe_del_mes(date(2026, 9, 1), False)['generado'].tzinfo)

    def test_con_facturacion_electronica_usa_los_anexos_de_sifen(self):
        Configuracion.establecer(CLAVE_FACTURACION_ELECTRONICA_ACTIVO, 'true')
        db.session.commit()
        http = self.app.test_client()
        with http.session_transaction() as sesion:
            sesion['_user_id'] = str(Usuario.query.filter_by(username='admin').first().id_usuario)
            sesion['_fresh'] = True

        texto, _paginas = self._texto_pdf(http.get(f'/contador/iva/pdf?mes={MES}'))

        self.assertIn('Facturas electrónicas aprobadas', texto)
        self.assertNotIn('Ventas con número de factura', texto)

    def test_la_pantalla_ofrece_las_descargas(self):
        html = self.http.get(f'/contador/iva?mes={MES}').get_data(as_text=True)

        self.assertIn(f'/contador/iva/pdf?mes={MES}', html)
        self.assertIn(f'/contador/iva/excel?mes={MES}', html)

    def test_excel_advierte_ambiente_de_prueba(self):
        from openpyxl import load_workbook
        from facturacion_electronica.models import FacturacionElectronicaConfig

        Configuracion.establecer(CLAVE_FACTURACION_ELECTRONICA_ACTIVO, 'true')
        FacturacionElectronicaConfig.obtener().ambiente = 'test'
        db.session.commit()
        respuesta = self.http.get(f'/contador/iva/excel?mes={MES}')
        libro = load_workbook(BytesIO(respuesta.data))
        self.assertTrue(any('sin validez fiscal' in str(fila[0])
                            for fila in libro['Resumen'].iter_rows(values_only=True)))
