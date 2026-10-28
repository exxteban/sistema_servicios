import io
from datetime import date, datetime

from openpyxl import Workbook

from app import create_app, db
from app.models import Caja, MovimientoCaja, SesionCaja, Venta
from gastronomia.importacion_models import GastronomiaImportacion, GastronomiaImportAlias
from gastronomia.models import GastronomiaPedido, GastronomiaPedidoPago, GastronomiaProducto
from gastronomia.services.importacion.dias import resolver_dias
from gastronomia.services.importacion.excel_lector import leer_excel
from gastronomia.services.importacion.fechas import parse_fecha
from gastronomia.services.importacion.productos_match import agrupar_alias
from test_gastronomia_reportes import _crear_productos, _csrf, _loguear


def _excel_desordenado() -> bytes:
    """Imita la planilla real: titulos corridos, bloques con 'Total', fechas escritas a mano."""
    libro = Workbook()
    ventas = libro.active
    ventas.title = 'venta por nota de pedido'
    ventas['B3'] = 'Ingresos'
    ventas.append([])
    ventas.append([None, 'Fecha ', 'cantidad', 'mercaderia', 'monto'])
    ventas.append([None, datetime(2026, 10, 14), 2, 'palitos', 16000, 8000])
    ventas.append([None, None, 1, 'flye dobke', 35000])
    ventas.append([None, None, 1, 'Flye Doble', 35000])
    ventas.append([None, None, None, '2 palito', 16000])
    ventas.append([None, None, None, 'total ', 102000])
    ventas.append([])
    ventas.append([None, None, 1, 'gus normal', 28000])
    ventas.append([None, None, 1, 'gus doble', 35000])
    ventas.append([None, None, 1, 'guss', 28000])
    ventas.append([None, None, None, 'Total', 91000])
    ventas.append([None, '07de cotubre', 1, 'hunk', 20000])
    ventas.append([None, None, 1, 'flye simple, 2flye doble', 98000])
    ventas.append([None, None, None, 'TOTAL', 100000])

    totales = libro.create_sheet('ingreso por dia')
    totales.append([])
    totales.append([None, datetime(2026, 10, 2), 91000])
    totales.append([None, datetime(2026, 10, 14), 102000])
    totales.append([None, 'Total', 193000])

    egresos = libro.create_sheet('egresos')
    egresos.append([None, 'Fecha', 'Monto', 'Razon Social'])
    egresos.append([None, datetime(2026, 10, 5), 92130, 'ticket'])

    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def test_fechas_escritas_a_mano():
    assert parse_fecha('07de cotubre', 2026) == date(2026, 10, 7)
    assert parse_fecha('15/0ct', 2026) == date(2026, 10, 15)
    assert parse_fecha('3 de noviembre', 2026) == date(2026, 11, 3)
    assert parse_fecha('12/10/25', 2026) == date(2025, 10, 12)
    assert parse_fecha('total', 2026) is None


def test_lector_detecta_hojas_bloques_y_fechas(tmp_path):
    ruta = tmp_path / 'ventas.xlsx'
    ruta.write_bytes(_excel_desordenado())

    lectura = leer_excel(str(ruta))

    tipos = {hoja['nombre']: hoja['tipo'] for hoja in lectura.hojas}
    assert tipos == {'venta por nota de pedido': 'ventas', 'ingreso por dia': 'totales', 'egresos': 'egresos'}
    assert [len(b.filas) for b in lectura.bloques] == [4, 3, 2]
    assert lectura.bloques[0].filas[3].cantidad == 2 and lectura.bloques[0].filas[3].nombre == 'palito'

    dias, sin_detalle = resolver_dias(lectura.bloques, lectura.totales_diarios)
    assert [d['fecha'] for d in dias] == ['2026-10-14', '2026-10-02', '2026-10-07']
    assert dias[1]['origen_fecha'] == 'totales'
    assert dias[2]['estado'] == 'revisar'  # suma 118.000 contra Total escrito 100.000
    assert sin_detalle == []


def test_agrupa_formas_de_escribir_en_productos():
    filas = [
        {'nombre': 'flye dobke', 'cantidad': 1, 'monto': 35000},
        {'nombre': 'Flye Doble', 'cantidad': 1, 'monto': 35000},
        {'nombre': 'flyedoble', 'cantidad': 1, 'monto': 35000},
        {'nombre': 'gus normal', 'cantidad': 1, 'monto': 28000},
        {'nombre': 'guss', 'cantidad': 1, 'monto': 28000},
        {'nombre': 'gus doble', 'cantidad': 1, 'monto': 35000},
        {'nombre': 'gus con panceta', 'cantidad': 1, 'monto': 28000},
        {'nombre': 'flye simple, 2flye doble', 'cantidad': 1, 'monto': 98000},
    ]
    grupos = {g['nombre']: g for g in agrupar_alias(filas, [], {})}

    assert grupos['Flye Doble']['unidades'] == 3
    assert grupos['Gus Simple']['unidades'] == 3
    assert grupos['Gus Doble']['unidades'] == 1
    assert grupos['Varios (importado)']['accion'] == 'varios'
    nota = next(a for a in grupos['Gus Simple']['alias'] if a['alias'] == 'gus con panceta')['nota']
    assert nota == 'con panceta'


def test_importa_deshace_y_aprende(tmp_path):
    app = create_app('testing')
    app.instance_path = str(tmp_path)
    client = app.test_client()
    cliente_id, _pizza_id, _bebida_id = _crear_productos(app, 'Resto Importa', 'resto_importa')
    with app.app_context():
        db.session.add(Caja(nombre='Caja Importa', ubicacion='Local'))
        db.session.commit()
    _loguear(client, app, 'resto_importa')
    csrf = _csrf(client.get('/gastronomia/importar-ventas').get_data(as_text=True))
    headers = {'X-CSRFToken': csrf}

    resp = client.post(
        '/api/gastronomia/importar-ventas/analizar',
        data={'archivo': (io.BytesIO(_excel_desordenado()), 'breaking.xlsx')},
        headers=headers,
        content_type='multipart/form-data',
    )
    assert resp.status_code == 200, resp.get_json()
    analisis = resp.get_json()['analisis']
    assert analisis['resumen'] == {'dias': 3, 'pedidos': 9, 'total': 311000.0}

    payload = {
        'token': analisis['token'],
        'nombre_archivo': 'breaking.xlsx',
        'dias': [{'indice': d['indice'], 'fecha': d['fecha'], 'incluir': True} for d in analisis['dias']],
    }
    resp = client.post('/api/gastronomia/importar-ventas/aplicar', json=payload, headers=headers)
    assert resp.status_code == 200, resp.get_json()
    importacion = resp.get_json()['importacion']
    assert importacion['pedidos'] == 9
    assert importacion['total'] == 311000.0
    assert 'Flye Doble' in importacion['productos_creados']

    with app.app_context():
        pagos = GastronomiaPedidoPago.query.filter_by(cliente_id=cliente_id).all()
        assert len(pagos) == 9
        assert sum(float(p.total_cobrado) for p in pagos) == 311000.0
        # Ya pagados: no deben aparecer como pendientes en caja, salon ni encargos.
        assert {p.estado for p in GastronomiaPedido.query.filter_by(cliente_id=cliente_id)} == {'cobrado'}
        sesiones = SesionCaja.query.filter(SesionCaja.observaciones.like('Importacion historica%')).all()
        assert len(sesiones) == 3 and all(s.estado == 'cerrada' for s in sesiones)
        assert SesionCaja.query.filter_by(estado='abierta').count() == 0
        assert sum(float(s.monto_final_sistema) for s in sesiones) == 311000.0
        assert GastronomiaImportAlias.query.filter_by(cliente_id=cliente_id, alias='flye dobke').count() == 1

    resumen = client.get('/api/gastronomia/reportes/resumen?desde=2026-10-14&hasta=2026-10-14').get_json()['resumen']
    assert resumen['ventas_total'] == 102000.0
    assert resumen['pedidos_cobrados'] == 4

    repetido = client.post('/api/gastronomia/importar-ventas/aplicar', json=payload, headers=headers)
    assert repetido.status_code == 400

    resp = client.post(f"/api/gastronomia/importar-ventas/{importacion['id_importacion']}/revertir", headers=headers)
    assert resp.status_code == 200, resp.get_json()
    assert resp.get_json()['importacion']['pedidos_borrados'] == 9
    with app.app_context():
        assert GastronomiaPedido.query.filter_by(cliente_id=cliente_id).count() == 0
        assert GastronomiaPedidoPago.query.filter_by(cliente_id=cliente_id).count() == 0
        assert Venta.query.filter(Venta.observaciones.like('%Importado de Excel%')).count() == 0
        assert MovimientoCaja.query.filter(MovimientoCaja.motivo.like('%(importado)%')).count() == 0
        assert SesionCaja.query.filter(SesionCaja.observaciones.like('Importacion historica%')).count() == 0
        assert db.session.get(GastronomiaImportacion, importacion['id_importacion']).estado == 'revertida'
        assert GastronomiaProducto.query.filter_by(cliente_id=cliente_id, nombre='Flye Doble').count() == 1

    # Lo aprendido hace que el proximo Excel use los productos ya creados.
    resp = client.post(
        '/api/gastronomia/importar-ventas/analizar',
        data={'archivo': (io.BytesIO(_excel_desordenado()), 'breaking.xlsx')},
        headers=headers,
        content_type='multipart/form-data',
    )
    grupos = {g['nombre']: g for g in resp.get_json()['analisis']['grupos']}
    assert grupos['Flye Doble']['accion'] == 'existente'
    assert grupos['Flye Doble']['origen'] == 'aprendido'
    # Lo que se mando a Varios tambien se recuerda, y Varios no se ofrece como plato del menu.
    assert grupos['Varios (importado)']['origen'] == 'aprendido'
    assert all(p['nombre'] != 'Varios (importado)' for p in resp.get_json()['analisis']['menu'])


def test_importar_ventas_requiere_sesion():
    app = create_app('testing')
    client = app.test_client()
    assert client.get('/api/gastronomia/importar-ventas').status_code in (302, 401)
    assert client.get('/gastronomia/importar-ventas').status_code in (302, 401)


def test_guarda_el_excel_aunque_instance_no_se_pueda_escribir(tmp_path):
    """En el servidor el servicio puede no tener permiso sobre instance/: se usa la carpeta temporal."""
    from werkzeug.datastructures import FileStorage
    from gastronomia.services.importacion.analisis import guardar_archivo, ruta_archivo

    app = create_app('testing')
    bloqueado = tmp_path / 'instance_es_un_archivo'
    bloqueado.write_text('x')
    app.instance_path = str(bloqueado)
    with app.app_context():
        token = guardar_archivo(FileStorage(io.BytesIO(_excel_desordenado()), 'ventas.xlsx'))
        ruta = ruta_archivo(token)
    assert not ruta.startswith(str(bloqueado))
    assert leer_excel(ruta).bloques
