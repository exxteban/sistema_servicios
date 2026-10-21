"""Gastos con factura en el crédito fiscal del resumen para el/la contador/a.

Un pago de gasto corriente con su factura (número, timbrado, RUC y tasa) suma
IVA al crédito fiscal del mes de la factura. Los retiros del dueño no. El
monto es siempre el del pago; cargar la factura no toca caja ni el gasto.
"""
from datetime import date
from decimal import Decimal

from app import db
from app.models import Configuracion, GastoComprobanteFiscal, GastoCorriente, PagoGastoCorriente, Usuario
from facturacion_electronica import CLAVE_FACTURACION_ELECTRONICA_ACTIVO

from test_facturacion_electronica_proveedor_base import BaseProveedorApi

MES = '2026-09'
PRIMERO = date(2026, 9, 1)


class TestGastosCreditoFiscal(BaseProveedorApi):
    def setUp(self):
        super().setUp()
        from facturacion_electronica.services.validacion import digito_verificador_ruc

        Configuracion.establecer(CLAVE_FACTURACION_ELECTRONICA_ACTIVO, 'false')
        db.session.commit()
        self.ruc = f'80012345-{digito_verificador_ruc("80012345")}'
        self.gasto = GastoCorriente(nombre='ANDE', categoria='servicios', monto_estimado=Decimal('110000'),
                                    dia_vencimiento=10, activo=True)
        db.session.add(self.gasto)
        db.session.commit()
        self.http = self.app.test_client()
        with self.http.session_transaction() as sesion:
            sesion['_user_id'] = str(Usuario.query.filter_by(username='admin').first().id_usuario)
            sesion['_fresh'] = True

    def _pago(self, monto=110000, estado='pagado', naturaleza='gasto_negocio', fecha=date(2026, 9, 12)):
        pago = PagoGastoCorriente(id_gasto_corriente=self.gasto.id_gasto_corriente, periodo_anio=2026, periodo_mes=9,
                                  fecha_vencimiento=date(2026, 9, 10), fecha_pago=fecha, monto_estimado=monto,
                                  monto_pagado=monto, estado=estado, naturaleza_contable=naturaleza)
        db.session.add(pago)
        db.session.commit()
        return pago

    def _factura(self, pago, **cambios):
        datos = {'numero': '001-002-345', 'timbrado': '15151515', 'ruc': self.ruc, 'razon_social': 'ANDE',
                 'fecha': '2026-09-12', 'tasa_iva': '10'}
        datos.update(cambios)
        return self.http.post(f'/contador/gasto-pago/{pago.id_pago_gasto_corriente}/factura', data=datos)

    def test_factura_del_gasto_suma_al_credito_fiscal(self):
        from app.services.resumen_contadora import resumen_del_mes

        pago = self._pago()
        r = resumen_del_mes(PRIMERO, fe_activa=False)
        self.assertEqual(r['credito'], 0)
        self.assertEqual(r['gastos_sin_factura'], {'cantidad': 1, 'total': 110000.0})

        self.assertEqual(self._factura(pago).status_code, 302)
        comprobante = GastoComprobanteFiscal.query.one()
        self.assertEqual(comprobante.numero, '001-002-0000345')

        r = resumen_del_mes(PRIMERO, fe_activa=False)
        self.assertAlmostEqual(r['credito'], 10000)
        self.assertAlmostEqual(r['gastos']['base10'], 100000)
        self.assertEqual(r['gastos_sin_factura']['cantidad'], 0)
        db.session.refresh(pago)
        self.assertEqual(float(pago.monto_pagado), 110000)            # el pago no cambia

        pantalla = self.http.get(f'/contador/iva?mes={MES}').get_data(as_text=True)
        self.assertIn("abrir('gastos')", pantalla)
        detalle = self.http.get(f'/contador/iva/detalle?mes={MES}&fila=gastos').get_data(as_text=True)
        self.assertIn('001-002-0000345', detalle)
        self.assertIn('15151515', detalle)
        self.assertIn(self.ruc, detalle)
        excel = self.http.get(f'/contador/iva/detalle?mes={MES}&fila=gastos&formato=xlsx')
        self.assertEqual(excel.status_code, 200)

    def test_tasa_cinco_y_exenta(self):
        from app.services.gastos_credito_fiscal import totales

        self.assertAlmostEqual(totales(21000, 5)['iva5'], 1000)
        self.assertEqual(totales(5000, 0)['exenta'], 5000)
        self.assertEqual(totales(5000, 0)['iva10'], 0)

    def test_facturas_repetidas_avisan_sin_descartar_pagos_parciales(self):
        from app.services.resumen_contadora_informe import informe_del_mes

        self._factura(self._pago(monto=55000))
        self._factura(self._pago(monto=55000))
        informe = informe_del_mes(PRIMERO, False)
        self.assertEqual(informe['resumen']['credito'], 10000)
        self.assertTrue(any('carga duplicada' in aviso for aviso in informe['pendientes']))
        html = self.http.get(f'/contador/iva?mes={MES}').get_data(as_text=True)
        self.assertIn('pagos parciales', html)

    def test_excel_conserva_precision_para_sumar_comprobantes(self):
        from io import BytesIO
        from openpyxl import load_workbook
        from app.services.resumen_contadora import resumen_del_mes

        for numero in range(3):
            self._factura(self._pago(monto=5), numero=f'001-001-{numero + 1:07d}')
        respuesta = self.http.get(f'/contador/iva/excel?mes={MES}')
        libro = load_workbook(BytesIO(respuesta.data))
        filas = list(libro['Gastos con factura'].iter_rows(values_only=True))
        columna_iva = filas[3].index('IVA 10%')
        self.assertAlmostEqual(sum(f[columna_iva] for f in filas[4:-1]), filas[-1][columna_iva])
        self.assertAlmostEqual(filas[-1][columna_iva], resumen_del_mes(PRIMERO, False)['credito'])
        resumen = list(libro['Resumen'].iter_rows(values_only=True))
        credito = next(f[2] for f in resumen if f[0] == 'Crédito fiscal')
        self.assertAlmostEqual(credito, filas[-1][columna_iva])

    def test_validaciones(self):
        pago = self._pago()
        self._factura(pago, timbrado='123')
        self._factura(pago, ruc='80012345-' + str((int(self.ruc[-1]) + 1) % 10))
        self._factura(pago, ruc='80012345')
        for ruc in ('-', 'abc-', '80012345-00'):
            self._factura(pago, ruc=ruc)
        self.assertEqual(GastoComprobanteFiscal.query.count(), 0)

        retiro = self._pago(naturaleza='retiro_dueno')
        self._factura(retiro)
        anulado = self._pago(estado='anulado')
        self._factura(anulado)
        self.assertEqual(GastoComprobanteFiscal.query.count(), 0)

        self._factura(pago)
        self.assertEqual(GastoComprobanteFiscal.query.count(), 1)
        self._factura(pago, numero='')
        self.assertEqual(GastoComprobanteFiscal.query.count(), 0)

    def test_cuenta_en_el_mes_de_la_factura_y_no_si_se_anula_el_pago(self):
        from app.services.resumen_contadora import resumen_del_mes

        pago = self._pago()
        self._factura(pago, fecha='2026-10-03')
        self.assertEqual(resumen_del_mes(PRIMERO, fe_activa=False)['gastos']['cantidad'], 0)
        self.assertEqual(resumen_del_mes(date(2026, 10, 1), fe_activa=False)['gastos']['cantidad'], 1)

        pago.estado = 'anulado'
        db.session.commit()
        self.assertEqual(resumen_del_mes(date(2026, 10, 1), fe_activa=False)['gastos']['cantidad'], 0)

    def test_pago_reclasificado_como_retiro_deja_de_sumar(self):
        from app.services.resumen_contadora import resumen_del_mes

        pago = self._pago()
        self._factura(pago)
        self.assertEqual(resumen_del_mes(PRIMERO, fe_activa=False)['gastos']['cantidad'], 1)

        pago.naturaleza_contable = 'retiro_dueno'
        db.session.commit()
        r = resumen_del_mes(PRIMERO, fe_activa=False)
        self.assertEqual(r['gastos']['cantidad'], 0)
        self.assertEqual(r['credito'], 0)

    def test_ficha_del_gasto_ofrece_cargar_la_factura(self):
        self._pago()
        self._pago(naturaleza='retiro_dueno', fecha=date(2026, 9, 13))

        html = self.http.get(f'/gastos-corrientes/{self.gasto.id_gasto_corriente}').get_data(as_text=True)

        self.assertEqual(html.count('Cargar factura (IVA)'), 1)
