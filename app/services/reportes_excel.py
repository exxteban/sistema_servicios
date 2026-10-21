"""Armado común de los Excel: hoja con título, encabezado y formato de montos.

Portado de sistema_silvio_cel (sólo la parte genérica; sus reportes propios no
se trajeron). Lo usa el resumen para el/la contador/a.
"""
from datetime import datetime
from io import BytesIO

from flask import make_response

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_DISPONIBLE = True
except ModuleNotFoundError:
    Workbook = None
    Alignment = Border = Font = PatternFill = Side = None
    get_column_letter = None
    OPENPYXL_DISPONIBLE = False


EXCEL_MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def _es_entero(valor):
    """True si el valor numerico no tiene parte decimal real (p. ej. 712.0 -> True)."""
    if isinstance(valor, bool):
        return False
    if isinstance(valor, int):
        return True
    if isinstance(valor, float):
        return valor.is_integer()
    return False


def _respuesta_excel(wb, filename):
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    response = make_response(output.getvalue())
    response.headers['Content-Type'] = EXCEL_MIME
    response.headers['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def _respuesta_no_disponible():
    response = make_response('La exportación a Excel no está disponible en este servidor.', 503)
    response.headers['Content-Type'] = 'text/plain; charset=utf-8'
    return response


def _ajustar_columnas(ws):
    for col_idx in range(1, ws.max_column + 1):
        col_letter = get_column_letter(col_idx)
        max_length = 0
        for cell in ws[col_letter]:
            if cell.value is None:
                continue
            max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max(max_length + 2, 12), 55)


def _llenar_hoja(ws, titulo, subtitulo, headers, rows, *, money_cols=None, number_cols=None):
    last_col = max(len(headers), 1)
    last_col_letter = get_column_letter(last_col)
    ws.merge_cells(f'A1:{last_col_letter}1')
    ws['A1'] = titulo
    ws['A1'].font = Font(bold=True, size=15, color='0F172A')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')

    ws.merge_cells(f'A2:{last_col_letter}2')
    ws['A2'] = f'{subtitulo} | Generado: {datetime.now().strftime("%d/%m/%Y %H:%M")}'
    ws['A2'].font = Font(size=10, color='475569')
    ws['A2'].alignment = Alignment(horizontal='center', vertical='center')

    ws.append([])
    header_row = 4
    ws.append(headers)

    header_fill = PatternFill(fill_type='solid', fgColor='059669')
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for row in rows:
        ws.append(row)

    money_cols = set(money_cols or [])
    number_cols = set(number_cols or [])
    for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row, min_col=1, max_col=last_col):
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center')
            if cell.column in money_cols:
                # Montos en guaranies: siempre enteros, separador de miles, sin decimales
                cell.number_format = '#,##0'
                cell.alignment = Alignment(horizontal='right', vertical='center')
            elif cell.column in number_cols:
                # Cantidades: mostrar como entero limpio cuando no hay decimales reales
                # (evita el punto colgante "712." del formato '#,##0.##') y conserva
                # los decimales solo cuando el valor realmente los tiene.
                cell.number_format = '#,##0' if _es_entero(cell.value) else '#,##0.##'
                cell.alignment = Alignment(horizontal='right', vertical='center')

    if ws.max_row >= header_row:
        ws.auto_filter.ref = f'A{header_row}:{last_col_letter}{ws.max_row}'
    ws.freeze_panes = 'A5'
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 20
    _ajustar_columnas(ws)


def _crear_workbook(titulo, subtitulo, headers, rows, *, sheet='Reporte', money_cols=None, number_cols=None):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet[:31]
    _llenar_hoja(ws, titulo, subtitulo, headers, rows, money_cols=money_cols, number_cols=number_cols)
    return wb


__all__ = ['EXCEL_MIME', 'OPENPYXL_DISPONIBLE', 'Workbook', '_crear_workbook', '_llenar_hoja',
           '_respuesta_excel', '_respuesta_no_disponible']
